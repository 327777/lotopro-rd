import os
import sys
import time
import ssl
import json
import urllib.request
from bs4 import BeautifulSoup
from datetime import datetime, timezone, timedelta
from motor_jugada_maestra import calcular_jugada_maestra

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

# Zona Horaria Fija de República Dominicana (UTC-4) sin importar dónde corra el servidor (Londres, EE.UU., etc.)
TZ_RD = timezone(timedelta(hours=-4))

def ahora_rd():
    return datetime.now(TZ_RD)

DIR_ACTUAL = os.path.dirname(os.path.abspath(__file__))
ARCHIVO_JSON = os.path.join(DIR_ACTUAL, "sorteos_hoy.json")
ARCHIVO_JS = os.path.join(DIR_ACTUAL, "datos_en_vivo.js")
ARCHIVO_HIST = os.path.join(DIR_ACTUAL, "historial_loterias.json")
ARCHIVO_HIST_JS = os.path.join(DIR_ACTUAL, "datos_historial.js")

# Mapeo exhaustivo de nombres web a IDs internos.
# Las búsquedas se ordenan de mayor longitud a menor para evitar colisiones (ej: 'la suerte 6pm' antes de 'la suerte').
MAPA_LOTERIAS = {
    # Anguila
    "anguilla 10am": "anguila_10am",
    "anguila 10am": "anguila_10am",
    "anguilla 10:00": "anguila_10am",
    "anguilla 1pm": "anguila_1pm",
    "anguila 1pm": "anguila_1pm",
    "anguilla 1:00": "anguila_1pm",
    "anguilla 6pm": "anguila_6pm",
    "anguila 6pm": "anguila_6pm",
    "anguilla 6:00": "anguila_6pm",
    "anguilla 9pm": "anguila_9pm",
    "anguila 9pm": "anguila_9pm",
    "anguilla 9:00": "anguila_9pm",
    # Primera
    "la primera noche": "primera_noche",
    "primera noche": "primera_noche",
    "la primera 8pm": "primera_noche",
    "la primera 8:00": "primera_noche",
    "la primera dia": "la_primera_dia",
    "la primera 12pm": "la_primera_dia",
    "la primera 12:00": "la_primera_dia",
    "la primera": "la_primera_dia",
    # Suerte
    "la suerte 6pm": "la_suerte_tarde",
    "la suerte 6:00": "la_suerte_tarde",
    "suerte 6pm": "la_suerte_tarde",
    "la suerte tarde": "la_suerte_tarde",
    "la suerte 12:30": "la_suerte_dia",
    "la suerte 12pm": "la_suerte_dia",
    "la suerte dia": "la_suerte_dia",
    "la suerte": "la_suerte_dia",
    # Real
    "real noche": "real_noche",
    "real 9pm": "real_noche",
    "real 9:00": "real_noche",
    "quiniela real": "quiniela_real",
    "real 1pm": "quiniela_real",
    "real 1:00": "quiniela_real",
    "real": "quiniela_real",
    # New Jersey
    "new jersey tarde": "new_jersey_dia",
    "new jersey 1pm": "new_jersey_dia",
    "new jersey 1:00": "new_jersey_dia",
    "new jersey": "new_jersey_dia",
    # Florida
    "florida noche": "florida_noche",
    "florida 10pm": "florida_noche",
    "florida 10:00": "florida_noche",
    "florida tarde": "florida_dia",
    "florida 1:30": "florida_dia",
    "florida 1pm": "florida_dia",
    "florida": "florida_dia",
    # Lotedom
    "lotedom": "lotedom",
    # Nacional / Gana Más
    "nacional gana más": "gana_mas",
    "nacional gana mas": "gana_mas",
    "gana más": "gana_mas",
    "gana mas": "gana_mas",
    "gana mas 2:30": "gana_mas",
    "nacional noche": "loteria_nacional",
    "loteria nacional": "loteria_nacional",
    "nacional 9pm": "loteria_nacional",
    # New York
    "new york noche": "ny_noche",
    "new york 10:30": "ny_noche",
    "new york tarde": "ny_tarde",
    "new york 2:30": "ny_tarde",
    # Loteka
    "quiniela loteka": "quiniela_loteka",
    "loteka 7:55": "quiniela_loteka",
    "loteka": "quiniela_loteka",
    # Leidsa
    "quiniela leidsa": "quiniela_leidsa",
    "leidsa 8:55": "quiniela_leidsa",
    "leidsa": "quiniela_leidsa"
}

# URLs individuales por si alguna lotería requiere consulta directa de respaldo
URLS_RESPALDO = {
    "anguila_10am": ("https://enloteria.com/resultados-anguilla-10am", "Anguilla 10AM"),
    "la_primera_dia": ("https://enloteria.com/resultados-la-primera", "La Primera"),
    "la_suerte_dia": ("https://enloteria.com/resultados-la-suerte", "La Suerte"),
    "quiniela_real": ("https://enloteria.com/resultados-real", "Real"),
    "anguila_1pm": ("https://enloteria.com/resultados-anguilla-1pm", "Anguilla 1PM"),
    "new_jersey_dia": ("https://enloteria.com/resultados-new-jersey-tarde", "New Jersey Tarde"),
    "florida_dia": ("https://enloteria.com/resultados-florida-tarde", "Florida Tarde"),
    "lotedom": ("https://enloteria.com/resultados-lotedom", "LoteDom"),
    "gana_mas": ("https://enloteria.com/resultados-gana-mas", "Nacional Gana"),
    "ny_tarde": ("https://enloteria.com/resultados-new-york-tarde", "New York Tarde"),
    "la_suerte_tarde": ("https://enloteria.com/resultados-la-suerte-6pm", "La Suerte 6PM"),
    "anguila_6pm": ("https://enloteria.com/resultados-anguilla-6pm", "Anguilla 6PM"),
    "quiniela_loteka": ("https://enloteria.com/resultados-loteka", "Loteka"),
    "primera_noche": ("https://enloteria.com/resultados-la-primera-noche", "La Primera Noche"),
    "quiniela_leidsa": ("https://enloteria.com/resultados-leidsa", "Leidsa"),
    "loteria_nacional": ("https://enloteria.com/resultados-nacional-noche", "Nacional Noche"),
    "real_noche": ("https://enloteria.com/resultados-real-noche", "Real Noche"),
    "anguila_9pm": ("https://enloteria.com/resultados-anguilla-9pm", "Anguilla 9PM"),
    "florida_noche": ("https://enloteria.com/resultados-florida-noche", "Florida Noche"),
    "ny_noche": ("https://enloteria.com/resultados-new-york-noche", "New York Noche")
}

def hora_ha_pasado(hora_str):
    try:
        ahora = ahora_rd()
        t = datetime.strptime(hora_str.strip(), "%I:%M %p")
        # El sorteo ya debió salir si la hora actual en RD es mayor o igual
        return (ahora.hour * 60 + ahora.minute) >= (t.hour * 60 + t.minute)
    except Exception:
        return True

def sincronizar_excel_recientes(sid, nombre_sorteo, fecha_dmy, premios):
    archivo_excel = os.path.join(DIR_ACTUAL, "Recientes.xlsx")
    if not os.path.exists(archivo_excel):
        return
    try:
        import openpyxl
        wb = openpyxl.load_workbook(archivo_excel)
        hoja_match = None
        nom_clean = nombre_sorteo.lower().replace(":", ".").replace(" ", "")
        for h in wb.sheetnames:
            h_clean = h.lower().replace(":", ".").replace(" ", "")
            if h_clean == nom_clean:
                hoja_match = h
                break
        if not hoja_match:
            for h in wb.sheetnames:
                h_clean = h.lower().replace(":", ".").replace(" ", "")
                if h_clean in nom_clean or nom_clean in h_clean:
                    hoja_match = h
                    break
        if hoja_match:
            ws = wb[hoja_match]
            if ws.max_row >= 2 and str(ws.cell(row=2, column=1).value).strip() == fecha_dmy:
                ws.cell(row=2, column=2, value=premios[0])
                ws.cell(row=2, column=3, value=premios[1])
                ws.cell(row=2, column=4, value=premios[2])
            else:
                ws.insert_rows(2)
                ws.cell(row=2, column=1, value=fecha_dmy)
                ws.cell(row=2, column=2, value=premios[0])
                ws.cell(row=2, column=3, value=premios[1])
                ws.cell(row=2, column=4, value=premios[2])
            wb.save(archivo_excel)
            print(f"📊 Recientes.xlsx sincronizado: {hoja_match} -> {premios}")
    except Exception as e:
        print(f"Nota sincronizando Recientes.xlsx: {e}")

def actualizar_todo():
    hora_actual_str = ahora_rd().strftime('%I:%M:%S %p')
    print(f"[{hora_actual_str} RD] 📡 Consultando resultados oficiales en vivo...")
    
    # 1. Cargar datos locales
    with open(ARCHIVO_JSON, "r", encoding="utf-8") as f:
        datos = json.load(f)

    # 2. Verificar cambio de día con HORA DE REPÚBLICA DOMINICANA (UTC-4)
    ahora = ahora_rd()
    hoy_iso = ahora.strftime("%Y-%m-%d")
    hoy_dmy = ahora.strftime("%d/%m/%Y")
    hubo_cambios = False
    
    if datos.get("fecha_hoy") != hoy_iso:
        print(f"🌅 ¡Nuevo día detectado en RD ({hoy_iso})! Archivando sorteos de ayer...")
        sorteos_ayer = []
        for s in datos.get("sorteos_hoy", []):
            if s.get("premios"):
                sorteos_ayer.append({
                    "nombre": s["nombre"],
                    "hora": s["hora"],
                    "premios": s["premios"]
                })
        if len(sorteos_ayer) >= 15:
            datos["sorteos_ayer"] = sorteos_ayer
            
        datos["fecha_ayer"] = datos.get("fecha_hoy", "")
        datos["fecha_hoy"] = hoy_iso
        
        for s in datos.get("sorteos_hoy", []):
            s["estado"] = "proximo"
            s["premios"] = None
            if "alerta_cascada" in s:
                del s["alerta_cascada"]

        nueva_pareja = calcular_jugada_maestra()
        datos["jugada_maestra_fija"]["pareja_oficial"] = nueva_pareja
        print(f"🎯 Nueva Pareja Maestra Oficial: [ {nueva_pareja[0]} ] × [ {nueva_pareja[1]} ]")
        hubo_cambios = True

    hist = {}
    if os.path.exists(ARCHIVO_HIST):
        with open(ARCHIVO_HIST, "r", encoding="utf-8") as f:
            hist = json.load(f)

    # 2.5 Orden Oficial Cronológico Estricto (10:00 AM -> 10:30 PM)
    ORDEN_OFICIAL_IDS = [
        "anguila_10am",      # 10:00 AM
        "la_primera_dia",    # 12:00 PM
        "lotedom",           # 12:00 PM (Ubicado en su orden exacto del mediodía)
        "la_suerte_dia",     # 12:30 PM
        "quiniela_real",     # 1:00 PM
        "anguila_1pm",       # 1:00 PM
        "new_jersey_dia",    # 1:00 PM
        "florida_dia",       # 1:30 PM
        "gana_mas",          # 2:30 PM
        "ny_tarde",          # 2:30 PM
        "la_suerte_tarde",   # 6:00 PM
        "anguila_6pm",       # 6:00 PM
        "quiniela_loteka",   # 7:55 PM
        "primera_noche",     # 8:00 PM
        "quiniela_leidsa",   # 8:55 PM
        "loteria_nacional",  # 9:00 PM
        "real_noche",        # 9:00 PM
        "anguila_9pm",       # 9:00 PM
        "florida_noche",     # 10:00 PM
        "ny_noche"           # 10:30 PM
    ]
    ORDEN_OFICIAL_NOMBRES = [
        "Anguila 10:00 AM",
        "La Primera 12:00 PM",
        "LoteDom 12:00 PM",
        "La Suerte 12:30 PM",
        "Quiniela Real 1:00 PM",
        "Anguila 1:00 PM",
        "New Jersey 1:00 PM",
        "Florida 1:30 PM",
        "Gana Más 2:30 PM",
        "New York 2:30 PM",
        "La Suerte 6:00 PM",
        "Anguila 6:00 PM",
        "Quiniela Loteka 7:55 PM",
        "La Primera 8:00 PM",
        "Quiniela Leidsa 8:55 PM",
        "Lotería Nacional 9:00 PM",
        "Real 9:00 PM",
        "Anguila 9:00 PM",
        "Florida 10:00 PM",
        "New York 10:30 PM"
    ]

    if "sorteos_hoy" in datos:
        orden_actual_hoy = [s.get("id") for s in datos["sorteos_hoy"]]
        hoy_ordenado = sorted(
            datos["sorteos_hoy"],
            key=lambda s: ORDEN_OFICIAL_IDS.index(s.get("id")) if s.get("id") in ORDEN_OFICIAL_IDS else 999
        )
        if [s.get("id") for s in hoy_ordenado] != orden_actual_hoy:
            datos["sorteos_hoy"] = hoy_ordenado
            hubo_cambios = True

    if "sorteos_ayer" in datos:
        orden_actual_ayer = [s.get("nombre") for s in datos["sorteos_ayer"]]
        ayer_ordenado = sorted(
            datos["sorteos_ayer"],
            key=lambda s: ORDEN_OFICIAL_NOMBRES.index(s.get("nombre")) if s.get("nombre") in ORDEN_OFICIAL_NOMBRES else 999
        )
        if [s.get("nombre") for s in ayer_ordenado] != orden_actual_ayer:
            datos["sorteos_ayer"] = ayer_ordenado
            hubo_cambios = True

    # 3. Descargar resultados de enloteria.com
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    
    url = "https://enloteria.com/"
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
    try:
        req = urllib.request.Request(url, headers=headers)
        html = urllib.request.urlopen(req, context=ctx, timeout=12).read().decode("utf-8", errors="ignore")
    except Exception as e:
        print(f"⚠️ Error conectando a la fuente: {e}")
        return False

    # Cortar estrictamente antes del bloque de "Ayer" para no mezclar nunca sorteos del día anterior
    if 'aria-label="Ayer' in html:
        html = html.split('aria-label="Ayer')[0]

    soup = BeautifulSoup(html, "html.parser")

    # Ordenar claves de búsqueda por longitud descendente para evitar colisiones
    claves_ordenadas = sorted(MAPA_LOTERIAS.keys(), key=len, reverse=True)

    # 4. Procesar tarjetas de resultados de la página principal
    import re
    sorteos_vistos_hoy = set()
    tarjetas_hoy_detectadas = 0
    
    for card in soup.find_all("div", class_="result-card"):
        name_el = card.find(class_="lottery-name")
        if not name_el:
            continue
        nom_web = name_el.get_text(strip=True).lower()
        
        # Extraer fecha de la tarjeta y validar el número exacto del día (inmune a 2026, 2027, 2028...)
        date_el = card.find(class_="result-date")
        date_txt = date_el.get_text(strip=True).lower() if date_el else ""
        m_dia = re.match(r"^\s*(\d+)", date_txt)
        if not m_dia or int(m_dia.group(1)) != ahora.day:
            continue
            
        tarjetas_hoy_detectadas += 1

        # Encontrar ID correspondiente sin colisiones
        sid = None
        for k in claves_ordenadas:
            if k in nom_web:
                sid = MAPA_LOTERIAS[k]
                break
        if not sid or sid in sorteos_vistos_hoy:
            continue

        # Extraer bolos
        balls = card.find_all(class_="result-number")
        nums = [int(b.get_text(strip=True)) for b in balls if b.get_text(strip=True).isdigit()]
        if len(nums) < 3:
            continue
        nums = nums[:3]

        sorteos_vistos_hoy.add(sid)

        # Buscar sorteo en sorteos_hoy
        for s in datos.get("sorteos_hoy", []):
            if s.get("id") == sid and hora_ha_pasado(s.get("hora", "")):
                if s["estado"] != "finalizado" or s.get("premios") != nums:
                    print(f"🎉 ¡NUEVO RESULTADO OFICIAL DETECTADO: {s['nombre']} -> {nums}!")
                    s["estado"] = "finalizado"
                    s["premios"] = nums
                    
                    # Actualizar historial
                    if sid not in hist:
                        hist[sid] = []
                    if hist[sid] and hist[sid][0].get("fecha") == hoy_dmy:
                        hist[sid][0]["premios"] = nums
                    else:
                        hist[sid].insert(0, {"fecha": hoy_dmy, "premios": nums})
                        
                    hubo_cambios = True
                    sincronizar_excel_recientes(sid, s.get("nombre", sid), hoy_dmy, nums)

    # 5. Auto-limpieza de seguridad: cualquier sorteo cuya hora aún no ha llegado hoy en RD
    # o que aún no tiene bolos publicados hoy en la fuente oficial se mantiene en "proximo"
    for s in datos.get("sorteos_hoy", []):
        sid_clean = s.get("id")
        sin_bolos_hoy = (tarjetas_hoy_detectadas >= 10 and sid_clean not in sorteos_vistos_hoy)
        if not hora_ha_pasado(s.get("hora", "")) or sin_bolos_hoy:
            if s.get("estado") != "proximo" or s.get("premios") is not None:
                s["estado"] = "proximo"
                s["premios"] = None
                hubo_cambios = True
            if sid_clean in hist and hist[sid_clean] and hist[sid_clean][0].get("fecha") == hoy_dmy:
                hist[sid_clean].pop(0)
                hubo_cambios = True

    if hubo_cambios:
        datos["actualizado_a_las"] = ahora.strftime("%I:%M %p")
        with open(ARCHIVO_JSON, "w", encoding="utf-8") as f:
            json.dump(datos, f, indent=2, ensure_ascii=False)
        with open(ARCHIVO_JS, "w", encoding="utf-8") as f:
            f.write("window.DATOS_LOTOPRO = " + json.dumps(datos, indent=2, ensure_ascii=False) + ";\n")
        with open(ARCHIVO_HIST, "w", encoding="utf-8") as f:
            json.dump(hist, f, indent=2, ensure_ascii=False)
        with open(ARCHIVO_HIST_JS, "w", encoding="utf-8") as f:
            f.write("window.HISTORIAL_LOTERIAS = " + json.dumps(hist, ensure_ascii=False) + ";\n")
        print("✅ Base de datos e historiales sincronizados al 100%.")
    else:
        print("✓ Pizarra al día. Sin sorteos nuevos por el momento.")

    return hubo_cambios

if __name__ == "__main__":
    print("==========================================================")
    print("   LOTOPRO RD — ACTUALIZADOR AUTÓNOMO 24/7 (ZONA HORARIA RD)")
    print("==========================================================")
    if "--bucle" in sys.argv:
        print("🚀 Modo Autónomo Continuo Activo (Rastrea cada 3 minutos)...")
        print("Presiona Ctrl + C para detener.\n")
        while True:
            try:
                actualizar_todo()
            except Exception as e:
                print(f"Error en ciclo: {e}")
            time.sleep(180)
    else:
        actualizar_todo()
